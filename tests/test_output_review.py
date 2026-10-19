import contextlib
import io
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from smetrics.filler import fill_back
from smetrics.models import PostMetric
from smetrics.writer import write_excel, print_summary


class OutputReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'accounts.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['账号名', '主页链接', 'Platform', 'Fans', 'Post', 'Post Date',
                   'Post Link', 'Views', 'Likes', 'Comments'])
        ws.append(['test', 'https://example.com/test', 'YouTube', 100, 'old post',
                   '2020-01-01', 'https://example.com/old', 999, 88, 77])
        ws['G2'].hyperlink = 'https://example.com/old'
        ws.append(['untouched', 'https://example.com/other', 'Instagram', 10, 'manual',
                   '2020-01-01', 'https://example.com/manual', 1, 2, 3])
        wb.save(self.source)
        wb.close()
        self.before = self.source.read_bytes()

    def metric(self, status='OK', **kwargs):
        return PostMetric(2, 'test', 'youtube', 'https://example.com/test',
                          target_date='2026-10-18', status=status, **kwargs)

    def fill(self, results):
        path, notes = fill_back(str(self.source), results, str(self.root/'output'), ['2026-10-18'])
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(self.source.read_bytes(), self.before)
        return wb, notes

    def test_no_post_or_failure_clears_stale_metrics_only_for_processed_row(self):
        for status in ['NO_POST', 'FAILED']:
            with self.subTest(status=status):
                wb, _ = self.fill([self.metric(status, error='limited')])
                ws = wb.active
                self.assertTrue(all(ws.cell(2, c).value is None for c in [4, 7, 8, 9, 10]))
                self.assertIsNone(ws['G2'].hyperlink)
                self.assertEqual(ws['I3'].value, 2)
                self.assertEqual(ws['E3'].value, 'manual')

    def test_partial_failure_remains_visible_in_filled_sheet(self):
        wb, notes = self.fill([self.metric(post_title='good', like_count=5),
                              self.metric('FAILED', error='video missing')])
        self.assertIn('部分内容抓取失败', wb.active['E2'].value)
        self.assertEqual(wb.active['I2'].value, 5)
        self.assertTrue(any('失败 1 行' in note for note in notes))

    def test_social_text_is_not_excel_formula(self):
        rows = [self.metric(post_title='=1+2')]
        target = self.root/'result.xlsx'
        write_excel(rows, str(target), {})
        wb = load_workbook(target, data_only=False)
        self.addCleanup(wb.close)
        self.assertEqual(wb['采集结果']['G2'].value, '=1+2')
        self.assertEqual(wb['采集结果']['G2'].data_type, 's')
        filled, _ = self.fill(rows)
        self.assertEqual(filled.active['E2'].value, '=1+2')
        self.assertEqual(filled.active['E2'].data_type, 's')

    def test_summary_counts_accounts_not_names_or_post_rows(self):
        rows = [self.metric(), PostMetric(3, 'test', 'instagram', 'url'),
                PostMetric(4, 'failed', 'instagram', 'url', status='FAILED'),
                PostMetric(4, 'failed', 'instagram', 'url', status='FAILED')]
        with contextlib.redirect_stdout(io.StringIO()) as out:
            print_summary(rows, [], {})
        self.assertIn('当天有发布：2 个账号', out.getvalue())
        self.assertIn('抓取失败　：1 个账号', out.getvalue())
