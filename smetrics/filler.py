"""把采集结果回填进用户原始表格的对应列。

原则：绝不改用户的原文件。复制原表后，刷新本次采集行的目标列。
未采集行和目标列之外的内容保持原样。
"""

import os
import shutil
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import STATUS_OK, STATUS_NO_POST, STATUS_FAILED, date_label as _date_label

# 目标列表头别名 → 内部字段。中英文都认，大小写和空格不敏感。
TARGET_COLUMNS = {
    "platform": ["platform", "平台", "渠道"],
    "follower": ["fans", "followers", "粉丝", "粉丝数", "订阅数", "账号粉丝数量"],
    "post_text": ["post", "帖子", "帖子文本", "内容", "标题", "作品标题"],
    "post_date": ["postdate", "date", "发帖日期", "发布日期", "发布时间"],
    "post_link": ["postlink", "link", "帖子链接", "作品链接", "视频链接"],
    "views": ["views", "view", "观看", "播放", "播放量", "观看数量", "观看数量（如有）"],
    "likes": ["likes", "like", "点赞", "点赞数", "点赞数量"],
    "comments": ["comments", "comment", "评论", "评论数", "评论数量"],
}

NO_POST_TEXT = "（当天无发布）"
NOT_SUPPORTED_TEXT = "（本次未采集）"
NONE_FILL = PatternFill("solid", fgColor="FFF2CC")
FAIL_FILL = PatternFill("solid", fgColor="FCE4E4")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")


def fill_back(source_path: str, results: list, out_dir: str,
              target_dates: list, touched_rows: set = None) -> tuple:
    """返回 (输出文件路径, 提示列表)。"""
    os.makedirs(out_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = os.path.splitext(os.path.basename(source_path))[0]
    label = _date_label(target_dates)
    extension = ".xlsm" if source_path.lower().endswith(".xlsm") else ".xlsx"
    out_path = os.path.join(out_dir, f"{base}-已回填-{label}-{stamp}{extension}")

    # 先物理复制，保证原文件一个字节都不动
    shutil.copy2(source_path, out_path)

    wb = load_workbook(out_path, keep_vba=extension == ".xlsm")
    ws = wb.active
    col_idx, header_row, notes = _locate_columns(ws)

    if not col_idx:
        wb.close()
        return out_path, ["原表里没找到 Fans/Post/Likes/Comments 这类目标列，已跳过回填"]

    # 同一行可能有多条作品（一天发了好几条），按行聚合后写进同一单元格，用换行分隔
    by_row = {}
    for r in results:
        by_row.setdefault(r.row_no, []).append(r)

    filled, no_post, failed = 0, 0, 0
    for row_no, rows in by_row.items():
        if row_no <= header_row:
            continue
        ok_rows = [r for r in rows if r.status == STATUS_OK]
        failed_rows = [r for r in rows if r.status == STATUS_FAILED]
        first = rows[0]

        # A filled input may contain last run's metrics. Never leave those next to
        # this run's NO_POST/FAILED status or stale links next to new post text.
        for field in TARGET_COLUMNS:
            if field != "platform":
                _write(ws, row_no, col_idx, field, None)
        _write(ws, row_no, col_idx, "platform", _platform_label(first.platform))
        if first.follower_count not in ("", None):
            _write(ws, row_no, col_idx, "follower", first.follower_count)

        if ok_rows:
            filled += 1
            title = _join(r.post_title for r in ok_rows)
            if failed_rows:
                failed += 1
                title += "\n（部分内容抓取失败，详见独立结果文件的待人工核对表）"
            _write(ws, row_no, col_idx, "post_text", title)
            _write(ws, row_no, col_idx, "post_date", _join(r.published_local for r in ok_rows))
            _write(ws, row_no, col_idx, "post_link", _join(r.post_url for r in ok_rows))
            _write(ws, row_no, col_idx, "views", _join_num(r.view_count for r in ok_rows))
            _write(ws, row_no, col_idx, "likes", _join_num(r.like_count for r in ok_rows))
            _write(ws, row_no, col_idx, "comments", _join_num(r.comment_count for r in ok_rows))
            _tint(ws, row_no, col_idx, FAIL_FILL if failed_rows else OK_FILL)
        elif failed_rows:
            failed += 1
            err = next(r.error for r in rows if r.status == STATUS_FAILED)
            # 单元格里只放简短原因，完整指引在「待人工核对」sheet，否则表格没法看
            brief = err.split("。")[0].split("解决办法")[0].strip()
            if len(brief) > 60:
                brief = brief[:60] + "…"
            _write(ws, row_no, col_idx, "post_text", f"（抓取失败：{brief} 详见独立结果文件的待人工核对表）")
            _write(ws, row_no, col_idx, "post_date", label)
            _tint(ws, row_no, col_idx, FAIL_FILL)
        else:
            no_post += 1
            hint = first.extra.get("提示") or ""
            _write(ws, row_no, col_idx, "post_text",
                   NO_POST_TEXT + (f" {hint}" if hint else ""))
            _write(ws, row_no, col_idx, "post_date", label)
            _tint(ws, row_no, col_idx, NONE_FILL)

    # 本次没跑到的行（比如 未选择的平台）留空不动，只在备注 sheet 里说明，避免覆盖手工填的内容
    _append_run_note(wb, target_dates, filled, no_post, failed, touched_rows, by_row)

    wb.save(out_path)
    wb.close()

    notes.append(f"回填完成：当天有发布 {filled} 行，无发布 {no_post} 行，失败 {failed} 行")
    covered = {r.platform for r in results}
    pending = [_platform_label(p) for p in ("youtube", "instagram") if p not in covered]
    if pending:
        notes.append(f"本次未采集的平台（{' / '.join(pending)}）整行留空未改动")
    return out_path, notes


def _locate_columns(ws) -> tuple:
    """在前 3 行里找表头，返回 (字段→列号, 表头所在行, 提示)。"""
    notes = []
    best, best_row = {}, 1
    for row_no in range(1, min(4, ws.max_row + 1)):
        found = {}
        for col in range(1, ws.max_column + 1):
            label = _norm(ws.cell(row=row_no, column=col).value)
            if not label:
                continue
            for field, aliases in TARGET_COLUMNS.items():
                if field in found:
                    continue
                if any(label == _norm(a) for a in aliases):
                    found[field] = col
                    break
        if len(found) > len(best):
            best, best_row = found, row_no

    if best:
        missing = [f for f in ("likes", "comments") if f not in best]
        if missing:
            notes.append(f"原表里没找到这些列，已跳过：{missing}")
    return best, best_row, notes


def _write(ws, row_no: int, col_idx: dict, field: str, value):
    col = col_idx.get(field)
    if not col:
        return
    cell = ws.cell(row=row_no, column=col)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Social captions beginning with '=' are text, never formulas.
    if field == "post_link":
        cell.hyperlink = None
    if isinstance(value, str) and "\n" in value:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if field == "post_link" and isinstance(value, str) and value.startswith("http") and "\n" not in value:
        cell.hyperlink = value
        cell.font = Font(color="0563C1", underline="single")


def _tint(ws, row_no: int, col_idx: dict, fill):
    col = col_idx.get("post_text") or col_idx.get("post_date")
    if col:
        ws.cell(row=row_no, column=col).fill = fill


def _join(values) -> str:
    return "\n".join(str(v) for v in values if str(v))


def _join_num(values) -> object:
    vals = list(values)
    if len(vals) == 1:
        return vals[0]
    return "\n".join(str(v) for v in vals)



def _platform_label(p: str) -> str:
    return {"youtube": "YouTube", "instagram": "Instagram"}.get(p, p)


def _norm(s) -> str:
    import re
    return re.sub(r"[\s_\-（）()]+", "", str(s or "")).lower()


def _append_run_note(wb, target_dates, filled, no_post, failed, touched_rows, by_row):
    name = "采集记录"
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        ws.append(["运行时间", "目标日期", "当天有发布", "当天无发布", "抓取失败", "回填行号"])
        for c in range(1, 7):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["F"].width = 46
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), _date_label(target_dates),
               filled, no_post, failed,
               ",".join(str(r) for r in sorted(by_row.keys()))])
