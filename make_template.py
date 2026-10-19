#!/usr/bin/env python3
"""生成账号表模板。直接跑：python3 make_template.py"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))

SAMPLES = [
    ("MrBeast", "https://www.youtube.com/@MrBeast", "YouTube"),
    ("Marques Brownlee", "https://www.youtube.com/@mkbhd", "YouTube"),
    ("Google", "https://www.youtube.com/channel/UCK8sQmJBp8GCxrOtXWBpyEA", "YouTube"),
    ("示例-Instagram", "https://www.instagram.com/natgeo/", "Instagram"),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "账号清单"
    ws.append(["账号名", "主页链接", "平台"])
    for c in range(1, 4):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
    for row in SAMPLES:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"

    tips = wb.create_sheet("填写说明")
    for line in [
        ["列", "说明"],
        ["账号名", "随便填，只用于输出时标识，可以是中文"],
        ["主页链接", "支持 youtube.com/@handle、/channel/UCxxx、/c/xxx、/user/xxx，"
                     "也能直接贴某个视频链接（会自动换成它所属频道）"],
        ["平台", "填 YouTube / Instagram 均可；填错也没关系，程序以链接为准"],
        ["", ""],
        ["注意", "列顺序可以打乱，程序按表头识别；没有表头时按 名称/链接/平台 三列顺序读"],
        ["注意", "Instagram 请填写主页链接；访问受限时需要登录态"],
    ]:
        tips.append(line)
    tips.column_dimensions["A"].width = 14
    tips.column_dimensions["B"].width = 92
    for c in range(1, 3):
        cell = tips.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")

    out = os.path.join(HERE, "accounts_template.xlsx")
    wb.save(out)
    print("模板已生成：" + out)


if __name__ == "__main__":
    main()
